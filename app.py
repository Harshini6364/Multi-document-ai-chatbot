import time
import tempfile
import re

import streamlit as st

from dotenv import load_dotenv

from langchain_groq import ChatGroq

from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader,
    UnstructuredPowerPointLoader
)

from langchain_core.documents import Document

from pptx import Presentation

from langchain_text_splitters import (
    RecursiveCharacterTextSplitter
)

from rank_bm25 import BM25Okapi


# -------------------------
# LOAD ENV
# -------------------------

load_dotenv()

# -------------------------
# LLM
# -------------------------

llm = ChatGroq(
    model="llama-3.3-70b-versatile",
    temperature=0
)

# -------------------------
# STREAMLIT
# -------------------------

st.set_page_config(
    page_title="Multi-Document AI Assistant",
    layout="wide"
)

st.title("📄 Multi-Document AI Assistant")

st.write(
    "Advanced Vector-less AI Assistant using BM25 + Groq"
)

# -------------------------
# SESSION STATE
# -------------------------

if "all_chunks" not in st.session_state:
    st.session_state.all_chunks = []

if "cache" not in st.session_state:
    st.session_state.cache = {}

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

# -------------------------
# TEXT PREPROCESSING
# -------------------------

def preprocess(text):

    text = text.lower()

    text = re.sub(
        r"[^a-zA-Z0-9 ]",
        " ",
        text
    )

    return text.split()

# -------------------------
# QUERY EXPANSION
# -------------------------

abbreviations = {
    "srs": "software requirement specification",
    "cn": "computer networks",
    "dbms": "database management system",
    "os": "operating system",
    "ai": "artificial intelligence",
    "ml": "machine learning",
    "nlp": "natural language processing"
}

# -------------------------
# SIDEBAR
# -------------------------

with st.sidebar:

    st.header("📂 Upload Documents")

    uploaded_files = st.file_uploader(
        "Upload Files",
        type=[
            "pdf",
            "txt",
            "pptx",
            "sql",
            "java",
            "py",
            "c",
            "cpp",
            "js",
            "html",
            "css"
        ],
        accept_multiple_files=True
    )

    st.markdown("---")

    st.subheader("📄 Uploaded Files")

    uploaded_sources = {
        chunk.metadata.get("source")
        for chunk in st.session_state.all_chunks
    }

    if uploaded_sources:

        for file in uploaded_sources:
            st.success(file)

    else:
        st.info("No files uploaded yet")

# -------------------------
# PROCESS FILES
# -------------------------

if uploaded_files:

    existing_sources = {
        chunk.metadata.get("source")
        for chunk in st.session_state.all_chunks
    }

    for uploaded_file in uploaded_files:

        if uploaded_file.name in existing_sources:
            continue

        ext = uploaded_file.name.split(".")[-1].lower()

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=f".{ext}"
        ) as tmp:

            tmp.write(uploaded_file.read())

            temp_path = tmp.name

        try:

            # -------------------------
            # PDF
            # -------------------------

            if ext == "pdf":

                loader = PyPDFLoader(temp_path)

                docs = loader.load()

            # -------------------------
            # PPTX
            # -------------------------

            elif ext == "pptx":

                prs = Presentation(temp_path)

                docs = []

                for slide_num, slide in enumerate(prs.slides):

                    slide_text = []

                    for shape in slide.shapes:

                        if hasattr(shape, "text"):

                            text = shape.text.strip()

                            if text:

                                slide_text.append(text)

                    combined_text = "\n".join(slide_text)

                    if combined_text.strip():

                        docs.append(
                            Document(
                                page_content=combined_text,
                                metadata={
                                    "page": slide_num + 1
                                }
                            )
                        )

            # -------------------------
            # TEXT FILES
            # -------------------------

            else:

                loader = TextLoader(
                    temp_path,
                    encoding="utf-8"
                )

                docs = loader.load()

        except Exception as e:

            st.error(
                f"Error loading "
                f"{uploaded_file.name}"
            )

            continue

        # -------------------------
        # TEXT SPLITTER
        # -------------------------

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=2000,
            chunk_overlap=400
        )

        chunks = splitter.split_documents(docs)

        for chunk in chunks:

            chunk.metadata["source"] = (
                uploaded_file.name
            )

            chunk.metadata["page"] = (
                chunk.metadata.get("page", 1)
            )

            st.session_state.all_chunks.append(
                chunk
            )

    st.success(
        "Documents processed successfully!"
    )

# -------------------------
# BM25
# -------------------------

bm25 = None

if st.session_state.all_chunks:

    chunk_texts = [
        chunk.page_content
        for chunk in st.session_state.all_chunks
    ]

    tokenized_chunks = [
        preprocess(text)
        for text in chunk_texts
    ]

    bm25 = BM25Okapi(tokenized_chunks)

# -------------------------
# OLD CHAT
# -------------------------

for chat in st.session_state.chat_history:

    with st.chat_message(chat["role"]):

        st.write(chat["message"])

        if chat["role"] == "assistant":

            st.markdown("---")

            st.markdown(
                f"⏱️ Response Time: "
                f"`{chat['response_time']} sec`"
            )

            if chat["citations"]:

                st.markdown(
                    "📚 Sources: "
                    + " | ".join(chat["citations"])
                )

# -------------------------
# USER QUERY
# -------------------------

query = st.chat_input(
    "Ask a question..."
)

if query and st.session_state.all_chunks:

    st.session_state.chat_history.append(
        {
            "role": "user",
            "message": query
        }
    )

    st.chat_message("user").write(query)

    normalized_query = query.lower().strip()

    expanded_query = normalized_query

    for short, full in abbreviations.items():

        if short in normalized_query:

            expanded_query += " " + full

    # -------------------------
    # CACHE
    # -------------------------

    if expanded_query in st.session_state.cache:

        start = time.time()

        cached = st.session_state.cache[
            expanded_query
        ]

        answer = cached["answer"]

        citations = cached["citations"]

        response_time = round(
            time.time() - start,
            4
        )

        with st.chat_message("assistant"):

            st.write(answer)

            st.markdown("---")

            st.markdown(
                f"⏱️ Response Time: "
                f"`{response_time} sec`"
            )

            if citations:

                st.markdown(
                    "📚 Sources: "
                    + " | ".join(citations)
                )

            st.success(
                "⚡ Response fetched from cache"
            )

        st.session_state.chat_history.append(
            {
                "role": "assistant",
                "message": answer,
                "response_time": response_time,
                "citations": citations,
                "cached": True
            }
        )

    else:

        start_time = time.time()

        # -------------------------
        # BM25 SEARCH
        # -------------------------

        scores = bm25.get_scores(
            preprocess(expanded_query)
        )

        scored_chunks = sorted(
            zip(
                st.session_state.all_chunks,
                scores
            ),
            key=lambda x: x[1],
            reverse=True
        )

        top_chunks = [
            (chunk, score)
            for chunk, score in scored_chunks
            if score > 0
        ][:8]

        # -------------------------
        # NO MATCH
        # -------------------------

        if not top_chunks:

            answer = (
                "The information is not available "
                "in the uploaded documents."
            )

            citations = []

        else:

            context = "\n\n".join([
                chunk.page_content
                for chunk, score in top_chunks
            ])

            prompt = f"""
You are a document question-answering assistant.

Rules:
1. Answer ONLY from the context.
2. Give partial answers if available.
3. Do NOT use outside knowledge.
4. Say 'The information is not available in the uploaded documents.'
only if no relevant content exists.

Context:
{context}

Question:
{query}
"""

            response = llm.invoke(prompt)

            answer = response.content.strip()

            citations = []

            shown = set()

            if (
                "information is not available"
                not in answer.lower()
            ):

                for chunk, score in top_chunks:

                    citation = (
                        f"{chunk.metadata.get('source')} "
                        f"- Page "
                        f"{chunk.metadata.get('page', 1)}"
                    )

                    if citation not in shown:

                        shown.add(citation)

                        citations.append(citation)

        response_time = round(
            time.time() - start_time,
            2
        )

        # -------------------------
        # SHOW RESPONSE
        # -------------------------

        with st.chat_message("assistant"):

            st.write(answer)

            st.markdown("---")

            st.markdown(
                f"⏱️ Response Time: "
                f"`{response_time} sec`"
            )

            if citations:

                st.markdown(
                    "📚 Sources: "
                    + " | ".join(citations)
                )

        # -------------------------
        # SAVE CHAT
        # -------------------------

        st.session_state.chat_history.append(
            {
                "role": "assistant",
                "message": answer,
                "response_time": response_time,
                "citations": citations,
                "cached": False
            }
        )

        # -------------------------
        # SAVE CACHE
        # -------------------------

        st.session_state.cache[
            expanded_query
        ] = {
            "answer": answer,
            "citations": citations
        }

# -------------------------
# NO FILE WARNING
# -------------------------

elif query and not st.session_state.all_chunks:

    st.warning(
        "Please upload at least one document."
    )